"""보고서 마크다운 -> .docx 변환.

범용 변환기가 아니라 이 보고서가 쓰는 문법만 다룬다:
  # ## ###  제목 / 본문 / **굵게** / `코드` / 파이프 표(<br> 포함) / ``` 코드블록
  > 인용 / * 목록 / --- 구분선
한글은 w:eastAsia 폰트를 함께 지정해야 Word 가 대체 폰트를 쓰지 않는다.
"""
import re, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

INK   = RGBColor(0x14, 0x18, 0x1E)
INK2  = RGBColor(0x52, 0x5C, 0x6B)
ACC   = RGBColor(0x1C, 0x5C, 0x4C)
RULEC = "D9DDE3"

def cfg(lang):
    if lang == "ko":
        return dict(body="맑은 고딕", east="맑은 고딕", head="맑은 고딕", mono="Consolas", size=10)
    return dict(body="Calibri", east="맑은 고딕", head="Calibri", mono="Consolas", size=10.5)

def set_font(run, name, east, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color is not None: run.font.color.rgb = color

def shade(cell, hexcolor):
    from docx.oxml import OxmlElement
    el = OxmlElement("w:shd"); el.set(qn("w:val"), "clear"); el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)

def borders(table):
    from docx.oxml import OxmlElement
    tbl = table._tbl
    pr = tbl.tblPr
    el = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), RULEC)
        el.append(e)
    pr.append(el)

TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\n]+?\*)")

def add_inline(par, text, C, size=None, base_bold=False):
    size = size if size is not None else C["size"]
    for part in TOKEN.split(text):
        if not part: continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = par.add_run(part[2:-2]); set_font(r, C["body"], C["east"], size, True, None, INK)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = par.add_run(part[1:-1]); set_font(r, C["mono"], C["east"], size - 0.5, base_bold, None, INK)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1]); set_font(r, C["body"], C["east"], size, base_bold, True, INK)
        else:
            r = par.add_run(part); set_font(r, C["body"], C["east"], size, base_bold, None, INK)

def split_row(line):
    s = line.strip()
    if s.startswith("|"): s = s[1:]
    if s.endswith("|"): s = s[:-1]
    return [c.strip() for c in s.split("|")]

def is_sep(line):
    return bool(re.fullmatch(r"\|?[\s:\-\|]+\|?", line.strip())) and "-" in line

def build(md_path, out_path, lang):
    C = cfg(lang)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.4)

    st = doc.styles["Normal"]
    st.font.name = C["body"]; st.font.size = Pt(C["size"])
    st.element.rPr.rFonts.set(qn("w:eastAsia"), C["east"])
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.line_spacing = 1.22

    lines = md_path.read_text().split("\n")
    i, n = 0, len(lines)
    while i < n:
        ln = lines[i]; s = ln.strip()

        if not s:
            i += 1; continue

        if s == "---":
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4)
            from docx.oxml import OxmlElement
            pb = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
            b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6")
            b.set(qn("w:space"), "1"); b.set(qn("w:color"), RULEC)
            pb.append(b); p._p.get_or_add_pPr().append(pb)
            i += 1; continue

        if s.startswith("```"):
            buf = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(9)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run("\n".join(buf))
            set_font(r, C["mono"], C["east"], C["size"] - 1, None, None, INK2)
            continue

        if s.startswith("#"):
            lvl = len(s) - len(s.lstrip("#"))
            txt = s[lvl:].strip()
            p = doc.add_paragraph()
            sizes = {1: 21, 2: 15.5, 3: 12.5, 4: 11.5}
            p.paragraph_format.space_before = Pt({1: 0, 2: 20, 3: 14, 4: 10}[min(lvl, 4)])
            p.paragraph_format.space_after = Pt({1: 10, 2: 6, 3: 4, 4: 3}[min(lvl, 4)])
            p.paragraph_format.keep_with_next = True
            for part in TOKEN.split(txt):
                if not part: continue
                if part.startswith("`") and part.endswith("`") and len(part) > 2:
                    r = p.add_run(part[1:-1])
                    set_font(r, C["mono"], C["east"], sizes[min(lvl, 4)] - 1, True, None,
                             ACC if lvl <= 2 else INK)
                elif part.startswith("**") and part.endswith("**") and len(part) > 4:
                    r = p.add_run(part[2:-2])
                    set_font(r, C["head"], C["east"], sizes[min(lvl, 4)], True, None,
                             ACC if lvl <= 2 else INK)
                else:
                    r = p.add_run(part)
                    set_font(r, C["head"], C["east"], sizes[min(lvl, 4)], True, None,
                             ACC if lvl <= 2 else INK)
            i += 1; continue

        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(6)
            from docx.oxml import OxmlElement
            pb = OxmlElement("w:pBdr"); b = OxmlElement("w:left")
            b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "18")
            b.set(qn("w:space"), "8"); b.set(qn("w:color"), "1C5C4C")
            pb.append(b); p._p.get_or_add_pPr().append(pb)
            chunks = []
            while i < n and lines[i].strip().startswith(">"):
                chunks.append(lines[i].strip().lstrip(">").strip()); i += 1
            add_inline(p, " ".join(chunks), C)
            continue

        if s.startswith("|"):
            rows, header = [], None
            while i < n and lines[i].strip().startswith("|"):
                if is_sep(lines[i]):
                    header = rows.pop() if rows else None
                else:
                    rows.append(split_row(lines[i]))
                i += 1
            ncol = max([len(r) for r in ([header] if header else []) + rows] or [1])
            t = doc.add_table(rows=0, cols=ncol)
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            borders(t)
            def fill(cells, data, bold, fillcolor):
                for k in range(ncol):
                    cell = cells[k]
                    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                    cell.paragraphs[0].paragraph_format.space_before = Pt(2)
                    txt = data[k] if k < len(data) else ""
                    pieces = txt.split("<br>")
                    for j, piece in enumerate(pieces):
                        par = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
                        par.paragraph_format.space_after = Pt(1)
                        add_inline(par, piece, C, C["size"] - 0.5, base_bold=bold)
                    if fillcolor: shade(cell, fillcolor)
            if header: fill(t.add_row().cells, header, True, "F2F4F6")
            for r in rows: fill(t.add_row().cells, r, False, None)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        if re.match(r"^[*\-] ", s):
            while i < n and re.match(r"^[*\-] ", lines[i].strip()):
                item = lines[i].strip()[2:]
                i += 1
                while i < n and lines[i].startswith("  ") and lines[i].strip() \
                        and not re.match(r"^[*\-] ", lines[i].strip()):
                    item += " " + lines[i].strip(); i += 1
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.left_indent = Cm(0.65)
                add_inline(p, item, C)
            continue

        if re.match(r"^\d+\. ", s):
            while i < n and re.match(r"^\d+\. ", lines[i].strip()):
                item = re.sub(r"^\d+\.\s*", "", lines[i].strip()); i += 1
                while i < n and lines[i].startswith("  ") and lines[i].strip() \
                        and not re.match(r"^\d+\. ", lines[i].strip()):
                    item += " " + lines[i].strip(); i += 1
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.left_indent = Cm(0.65)
                add_inline(p, item, C)
            continue

        buf = [s]; i += 1
        while i < n and lines[i].strip() and not re.match(
                r"^(#|\||>|```|---$|[*\-] |\d+\. )", lines[i].strip()):
            buf.append(lines[i].strip()); i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(buf), C)

    doc.save(str(out_path))
    return out_path

if __name__ == "__main__":
    import pathlib
    build(pathlib.Path(sys.argv[1]), pathlib.Path(sys.argv[2]), sys.argv[3])
    print("saved", sys.argv[2])
